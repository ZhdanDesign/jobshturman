"""Acceptance for the hrs CLI: one MCP server with local app, collector and document tools."""
import asyncio
import json
import os
import subprocess
import sys

import httpx
from collector_fixtures import T0, mock_http, tg_page  # loaded in conftest.py
from docx import Document
from hrs_cli import paths
from hrs_cli.server import build
from hrs_collector.service import Collector
from hrs_collector.store import Store
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

POST = ("<b>Вакансия: Финансовый директор (CFO)</b><br/><br/><b>Компания:</b> Холдинг Пример<br/>"
        "<b>Локация:</b> Москва, гибрид<br/>Требования: 10+ лет. Резюме: cv@example.com")
LOCAL_TOOLS = {"status_report", "intake_save_draft", "resume_save", "vacancies_sync", "applications_list"}
COLLECTOR_TOOLS = {"list_sources", "add_telegram_channel", "run_collection", "search_vacancies"}


def resume_docx(path):
    doc = Document()
    doc.add_paragraph("Иван Пример, финансовый директор")
    doc.add_paragraph("2021–2024, Альфа: рост выручки на 25%")
    doc.save(path)
    return path


def text_of(result):
    return " ".join(getattr(c, "text", "") for c in result.content) if hasattr(result, "content") else str(result)


def test_one_server_collects_into_cards_without_a_second_process(tmp_path, monkeypatch):
    monkeypatch.setenv("HRSHTURMAN_HOME", str(tmp_path))
    monkeypatch.setenv("HRS_PACE_FACTOR", "0")
    monkeypatch.delenv("HRSHTURMAN_WORKSPACE", raising=False)
    collector = Collector(tmp_path / "collector", store=Store(tmp_path / "collector.sqlite3"),
                          http_factory=mock_http({"t.me/s/execjobs": lambda r: httpx.Response(
                              200, text=tg_page("execjobs", -42, [(5, T0, POST, False)], has_prev=False))}))
    server = build(paths.settings(), collector)

    async def go():
        tools = [t.name for t in await server.list_tools()]
        assert len(tools) == len(set(tools))
        assert LOCAL_TOOLS | COLLECTOR_TOOLS | {"document_import"} <= set(tools)
        await server.call_tool("add_telegram_channel", {"source_id": "TG-M", "username": "@execjobs"})
        synced = await server.call_tool("vacancies_sync", {"run": True, "lookback_days": 30})
        cards = await server.call_tool("applications_list", {})
        return str(synced), str(cards)

    synced, cards = asyncio.run(go())
    assert "service_unreachable" not in synced
    assert "Финансовый директор" in cards


def test_serve_over_stdio_imports_a_document(tmp_path):
    env = dict(os.environ, HRSHTURMAN_HOME=str(tmp_path / "home"), HRSHTURMAN_CANDIDATE="synthetic-1")
    env.pop("HRSHTURMAN_WORKSPACE", None)
    params = StdioServerParameters(command=sys.executable, args=["-m", "hrs_cli", "serve"], env=env)
    cv = resume_docx(tmp_path / "cv.docx")

    async def go():
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write) as s:
                await s.initialize()
                tools = {t.name for t in (await s.list_tools()).tools}
                imported = await s.call_tool("document_import", {"path": str(cv)})
                status = await s.call_tool("status_report", {})
                return tools, imported, status

    tools, imported, status = asyncio.run(go())
    assert LOCAL_TOOLS | COLLECTOR_TOOLS | {"document_import"} <= tools
    assert not imported.is_error and not status.is_error
    data = imported.structured_content or json.loads(imported.content[0].text)
    assert data["source_id"].startswith("file-") and "рост выручки на 25%" in data["markdown"]
    assert (tmp_path / "home" / "workspace" / "documents" / f"{data['source_id']}.md").is_file()


def run_hrs(*args, home):
    env = dict(os.environ, HRSHTURMAN_HOME=str(home))
    env.pop("HRSHTURMAN_WORKSPACE", None)
    env.pop("HRS_COLLECTOR_DATA", None)
    return subprocess.run([sys.executable, "-m", "hrs_cli", *args], env=env, capture_output=True, text=True, timeout=120, check=False)


def test_cli_keeps_everything_in_one_data_folder(tmp_path):
    home = tmp_path / "home"
    out = json.loads(run_hrs("paths", home=home).stdout)
    assert out["workspace"] == str(home / "workspace") and out["collector"] == str(home / "collector")

    added = run_hrs("collect", "add-telegram", "TG-X", "--username", "@execjobs", home=home)
    assert added.returncode == 0, added.stderr
    sources = run_hrs("collect", "sources", home=home)
    assert "TG-X" in sources.stdout and (home / "collector").is_dir()

    doc = run_hrs("doc", "import", str(resume_docx(tmp_path / "cv.docx")), home=home)
    assert doc.returncode == 0, doc.stderr
    assert json.loads(doc.stdout)["chars"] > 0

    bad = run_hrs("doc", "import", str(tmp_path / "old.doc"), home=home)
    assert bad.returncode == 1 and json.loads(bad.stdout)["error"] == "unsupported_format"
